#!/usr/bin/env python3
"""
depository_review_sheet.py -- turn a diff_depositories.py run into the reviewer's workbook.

Tabs: Read me | Conflicts (B) | NGC has more (A) | Fill (NGC only). The C rows
(AntCat cites more) are left out on purpose: they are almost always AntCat being
more complete, and runbook principle 0.3 says not to spend the reviewer on them.

  python3 bolton/depository_review_sheet.py --run <diff_depositories --out dir> \\
      --protonyms protonyms_species_typeinfo.jsonl [--hints hints.csv] --xlsx out.xlsx

--hints: optional CSV (ngc_code,antcat_code,hint). A row whose NGC-only and
AntCat-only codes include the pair gets the hint text.
"""
import argparse, csv, json, os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

DECISIONS = ['AntCat correct', 'Fix AntCat', 'Same institution: add [Also]', 'Unsure']
FONT = 'Arial'
YELLOW = PatternFill('solid', start_color='FFFF99')


def style_header(ws, widths):
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(name=FONT, bold=True)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = 'A2'


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--run', required=True)
    ap.add_argument('--protonyms', required=True)
    ap.add_argument('--hints')
    ap.add_argument('--xlsx', required=True)
    a = ap.parse_args()

    names = {}
    for l in open(a.protonyms, encoding='utf-8'):
        if l.strip():
            r = json.loads(l); names[str(r['protonym_id'])] = r.get('name') or ''
    hints = []
    if a.hints:
        hints = [(h['ngc_code'], h['antcat_code'], h['hint']) for h in csv.DictReader(open(a.hints, encoding='utf-8'))]
    metrics = json.load(open(os.path.join(a.run, 'metrics.json'), encoding='utf-8'))
    resid = list(csv.DictReader(open(os.path.join(a.run, 'type_depository_residual.csv'), encoding='utf-8')))
    fill = list(csv.DictReader(open(os.path.join(a.run, 'type_fill_candidates.csv'), encoding='utf-8')))

    wb = Workbook()
    rd = wb.active; rd.title = 'Read me'
    B = sorted((r for r in resid if r['priority'] == 'B_both_differ'), key=lambda r: names.get(r['protonym_id'], ''))
    A = sorted((r for r in resid if r['priority'] == 'A_ngc_has_extra'), key=lambda r: names.get(r['protonym_id'], ''))
    lines = [
        ('Type depositories: NGC 2026 vs AntCat', True),
        (f"Run {metrics.get('mode')} of bolton/diff_depositories.py on AntCat exports of 23 Oct 2026 and the NGC 2026 catalogue.", False),
        ('Depository codes were first matched through the [Also] notes in AntCat Institutions (including your recent edits), '
         'your crosswalk sheet, Bolton\'s own depository list, and your confirmed errata. What is left are real differences.', False),
        (f"Concordance {metrics['pct_final']}% of {metrics['comparable']:,} comparable records. "
         f"{metrics['residual']} records still differ; this workbook holds the {len(B) + len(A)} that need a decision, plus {len(fill)} fill candidates.", False),
        ('', False),
        ('Tabs', True),
        (f"Conflicts (B), {len(B)} rows: NGC and AntCat each name a depository the other lacks. Start here.", False),
        (f"NGC has more (A), {len(A)} rows: NGC names a depository AntCat lacks, usually a paratype depository.", False),
        (f"Fill, {len(fill)} rows: NGC has type information and AntCat has none.", False),
        ('Left out: records where only AntCat names more depositories. AntCat is almost always the more complete side there.', False),
        ('', False),
        ('What to fill in', True),
        ('Only the yellow columns: "Decision" (pick from the list) and "Note".', False),
        ('Decision values: ' + '; '.join(DECISIONS) + '.', False),
        ('"Same institution: add [Also]" means the two codes are one collection; add the [Also] note in AntCat Institutions and the next run clears the row.', False),
        ('"Hint" is a pointer from the comparison, not a conclusion.', False),
        ('', False),
        ('Example row (Conflicts tab)', True),
        ('Protonym 200222, Blepharidatta delabiei Brandão 2015. NGC only: UFGB. AntCat only: HCJG. '
         'Decision: Same institution: add [Also]. Note: HCJG is the Hymenoptera collection at UFG Jataí.', False),
    ]
    for i, (t, bold) in enumerate(lines, 1):
        c = rd.cell(row=i, column=1, value=t)
        c.font = Font(name=FONT, bold=bold, size=12 if i == 1 else 10)
        c.alignment = Alignment(wrap_text=True, vertical='top')
    rd.column_dimensions['A'].width = 120

    head = ['AntCat protonym', 'Name (AntCat)', 'Author', 'Year', 'NGC only', 'AntCat only',
            'NGC depository line', 'AntCat depository text', 'Hint', 'Decision', 'Note']
    widths = [14, 44, 18, 7, 16, 16, 60, 60, 36, 24, 30]
    dv_ref = f'"{",".join(DECISIONS)}"'
    for title, rows in (('Conflicts (B)', B), ('NGC has more (A)', A)):
        ws = wb.create_sheet(title)
        ws.append(head); style_header(ws, widths)
        dv = DataValidation(type='list', formula1=dv_ref, allow_blank=True); ws.add_data_validation(dv)
        for r in rows:
            no, ao = set(filter(None, r['ngc_only'].split('; '))), set(filter(None, r['antcat_only'].split('; ')))
            hint = '; '.join(h for n, c, h in hints if n in no and c in ao)
            ws.append([int(r['protonym_id']), names.get(r['protonym_id'], ''), r['author'], r['year'],
                       r['ngc_only'], r['antcat_only'], r['ngc_depository_line'], r['antcat_depository_text'],
                       hint, '', ''])
            n = ws.max_row
            ws.cell(row=n, column=1).hyperlink = r['antcat_url']
            for col in range(1, len(head) + 1):
                c = ws.cell(row=n, column=col)
                c.font = Font(name=FONT, size=10, color='0000EE' if col == 1 else None, underline='single' if col == 1 else None)
                c.alignment = Alignment(wrap_text=col in (7, 8, 9, 11), vertical='top')
            for col in (10, 11):
                ws.cell(row=n, column=col).fill = YELLOW
            dv.add(f'J{n}')

    ws = wb.create_sheet('Fill')
    fh = ['AntCat protonym', 'Name (AntCat)', 'Author', 'Year', 'NGC type block', 'Done (initials)', 'Note']
    ws.append(fh); style_header(ws, [14, 34, 18, 7, 110, 14, 30])
    for r in fill:
        ws.append([int(r['protonym_id']), r['antcat_name'], r['author'], r['year'], r['ngc_type_block'], '', ''])
        n = ws.max_row
        ws.cell(row=n, column=1).hyperlink = f"https://www.antcat.org/protonyms/{r['protonym_id']}"
        for col in range(1, len(fh) + 1):
            c = ws.cell(row=n, column=col)
            c.font = Font(name=FONT, size=10, color='0000EE' if col == 1 else None, underline='single' if col == 1 else None)
            c.alignment = Alignment(wrap_text=col in (5, 7), vertical='top')
        for col in (6, 7):
            ws.cell(row=n, column=col).fill = YELLOW
    wb.save(a.xlsx)
    print(f"wrote {a.xlsx}: Conflicts {len(B)}, NGC has more {len(A)}, Fill {len(fill)}")


if __name__ == '__main__':
    main()
